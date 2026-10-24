import io
import sys
import docx
import httpx
import pypdf

BASE_URL = "http://127.0.0.1:8000"

def generate_pdf_bytes(clause_text: str) -> bytes:
    escaped = clause_text.replace("(", "\\(").replace(")", "\\)")
    stream_content = f"BT\n/F1 12 Tf\n72 712 Td\n({escaped}) Tj\nET\n".encode("utf-8")
    stream_len = len(stream_content)
    
    pdf_template = (
        b"%PDF-1.4\n"
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        b"4 0 obj << /Length " + str(stream_len).encode("utf-8") + b" >>\nstream\n"
        + stream_content
        + b"endstream\nendobj\n"
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000244 00000 n \n"
        b"0000000418 00000 n \n"
        b"trailer << /Size 6 /Root 1 0 R >>\n"
        b"startxref\n500\n%%EOF"
    )
    return pdf_template

def generate_docx_bytes(title: str, clause_text: str) -> bytes:
    doc = docx.Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(clause_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def main():
    print(f"Connecting to live FastAPI at {BASE_URL}...")
    client = httpx.Client(base_url=BASE_URL, timeout=10.0)

    # 1. Health Endpoint
    print("\n[1] Testing GET /api/ai/health")
    res = client.get("/api/ai/health")
    assert res.status_code == 200, f"Health check failed: {res.status_code} {res.text}"
    health_data = res.json()
    print(" -> Health response:", health_data)
    assert health_data["status"] == "healthy"

    # 2. Extract Endpoints (TXT, DOCX, PDF)
    txt_content = (
        "Section 14.3. Indemnification.\n"
        "Customer shall defend, indemnify, and hold harmless Provider without limitation as to aggregate damages.\n"
        "Section 18.2. Termination for convenience upon 15 days notice without cause."
    )
    txt_bytes = txt_content.encode("utf-8")
    docx_bytes = generate_docx_bytes("Vendor Master Agreement", txt_content)
    pdf_bytes = generate_pdf_bytes("Section 14.3. Customer shall defend, indemnify, and hold harmless Provider without limitation.")

    print("\n[2] Testing POST /api/ai/extract with TXT, DOCX, and PDF")
    # TXT
    res_txt = client.post("/api/ai/extract", files={"file": ("sample.txt", txt_bytes, "text/plain")})
    assert res_txt.status_code == 200, f"Extract TXT failed: {res_txt.text}"
    txt_extract = res_txt.json()
    print(f" -> TXT: {txt_extract['file_type']}, {txt_extract['total_words']} words, pages: {txt_extract['total_pages']}")
    assert "Customer shall defend" in txt_extract["full_text"]

    # DOCX
    res_docx = client.post("/api/ai/extract", files={"file": ("sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res_docx.status_code == 200, f"Extract DOCX failed: {res_docx.text}"
    docx_extract = res_docx.json()
    print(f" -> DOCX: {docx_extract['file_type']}, {docx_extract['total_words']} words, pages: {docx_extract['total_pages']}")
    assert "Customer shall defend" in docx_extract["full_text"]

    # PDF
    res_pdf = client.post("/api/ai/extract", files={"file": ("sample.pdf", pdf_bytes, "application/pdf")})
    assert res_pdf.status_code == 200, f"Extract PDF failed: {res_pdf.text}"
    pdf_extract = res_pdf.json()
    print(f" -> PDF: {pdf_extract['file_type']}, {pdf_extract['total_words']} words, pages: {pdf_extract['total_pages']}")
    assert "Customer shall defend" in pdf_extract["full_text"]

    # 3. Summarize Endpoints
    print("\n[3] Testing POST /api/ai/summarize (Upload & JSON)")
    res_sum_file = client.post("/api/ai/summarize", files={"file": ("contract.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res_sum_file.status_code == 200
    sum_data = res_sum_file.json()
    print(f" -> Summarize File: Score={sum_data['compliance_score']}, Risks={sum_data['total_risks']}")
    assert len(sum_data["key_obligations"]) > 0
    assert len(sum_data["critical_dates"]) > 0

    res_sum_text = client.post("/api/ai/summarize/text", json={"text": txt_content, "filename": "inline.txt"})
    assert res_sum_text.status_code == 200
    assert res_sum_text.json()["compliance_score"] > 0

    # 4. Analyze Endpoints with Evidence Verification
    print("\n[4] Testing POST /api/ai/analyze and verifying evidence references")
    res_analysis = client.post("/api/ai/analyze", files={"file": ("master_agreement.txt", txt_bytes, "text/plain")})
    assert res_analysis.status_code == 200
    analysis_data = res_analysis.json()
    print(f" -> Overall Risk Level: {analysis_data['overall_risk_level']}")
    print(f" -> High: {analysis_data['summary']['risk_counts']['high']}, Med: {analysis_data['summary']['risk_counts']['medium']}")
    
    assert len(analysis_data["findings"]) > 0, "No findings returned!"
    for f in analysis_data["findings"]:
        print(f"   [Finding {f['id']}] {f['title']} (Severity: {f['severity']})")
        print(f"     Clause: {f['clause_reference']}")
        print(f"     Evidence: {f['evidence_snippet']}")
        # Verify that evidence snippet corresponds to the original text
        # If snippet contains a matched substring from txt_content
        assert any(word in txt_content for word in f['evidence_snippet'].split() if len(word) > 5), \
            f"Evidence snippet '{f['evidence_snippet']}' does not correspond to original text!"
    print(" -> All evidence references verified against original text!")

    # 5. Diff Endpoint
    print("\n[5] Testing POST /api/ai/diff")
    diff_payload = {
        "comparison_id": "comp-e2e-audit",
        "title": "Vendor MSA v1 vs v2",
        "base_text": "Section 1. Liability cap shall be equal to twelve months fees paid.\nSection 2. Governed by laws of California.",
        "target_text": "Section 1. Liability cap shall be $10,000 for Vendor, uncapped for Customer.\nSection 2. Governed by laws of New York with jury waiver.",
        "base_label": "v1.0-Standard",
        "target_label": "v2.0-Vendor-Markup"
    }
    res_diff = client.post("/api/ai/diff", json=diff_payload)
    assert res_diff.status_code == 200
    diff_data = res_diff.json()
    print(f" -> Diff changes detected: {diff_data['total_changes']}")
    for d in diff_data["diffs"]:
        print(f"   Clause: {d['clause_title']} | Type: {d['diff_type']} | Impact: {d['risk_impact']}")
        print(f"   Old: {d['old_text']}")
        print(f"   New: {d['new_text']}")

    print("\n>>> ALL AI SERVICE ENDPOINTS VERIFIED VIA REAL HTTP REQUESTS SUCCESSFULLY! <<<")

if __name__ == "__main__":
    main()
